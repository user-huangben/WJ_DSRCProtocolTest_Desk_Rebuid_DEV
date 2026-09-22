from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# 工具脚本位于 .cursor/skills/operation-manual-docx/scripts/；生成物仍写入解决方案目录。
WORKSPACE = Path(__file__).resolve().parents[4]
ROOT = WORKSPACE / 'WJ_DSRCProtocolTest_Desk_Net'
OUT = ROOT / '操作说明.docx'
ART = ROOT / 'artifacts' / 'manual'
SRC = ART / 'mainform.png'

def font(size, bold=False):
    candidates = [r'C:\Windows\Fonts\msyh.ttc', r'C:\Windows\Fonts\simhei.ttf']
    for p in candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size=size)
    return ImageFont.load_default()

def annotate(boxes, output):
    im = Image.open(SRC).convert('RGB')
    d = ImageDraw.Draw(im)
    f = font(18, True)
    for i, (box, label) in enumerate(boxes, 1):
        x1, y1, x2, y2 = box
        d.rectangle(box, outline=(220, 40, 40), width=4)
        d.ellipse((x1, y1, x1 + 30, y1 + 30), fill=(220, 40, 40))
        d.text((x1 + 8, y1 + 3), str(i), fill='white', font=f)
    im.save(output)

annotate([
    ((8, 25, 1280, 115), '一级产品 Tab 与标题区'),
    ((10, 116, 275, 780), '台发设置与状态'),
    ((295, 128, 560, 780), '帧解析与日志区'),
    ((655, 128, 1245, 345), '实时上下行帧区'),
    ((655, 350, 1245, 745), '帧操作与测试区'),
], ART / 'mainform_annotated.png')
annotate([
    ((10, 116, 275, 780), '端口、台发参数与更新选项'),
    ((655, 350, 1245, 745), '帧列表、选择器、测试与 BUF'),
], ART / 'mainform_regions.png')

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(text); r.bold = bold
    if color: r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for c, h in zip(t.rows[0].cells, headers):
        set_cell_text(c, h, True, (255,255,255)); shade(c, '1F4E78')
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for c, value in zip(cells, row):
            set_cell_text(c, value)
            if ridx % 2 == 1: shade(c, 'EAF2F8')
    doc.add_paragraph()
    return t

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.7); sec.right_margin = Inches(0.7)
styles = doc.styles
styles['Normal'].font.name = 'Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); styles['Normal'].font.size = Pt(10)
for name, size in [('Title', 24), ('Heading 1', 17), ('Heading 2', 14), ('Heading 3', 11)]:
    styles[name].font.name = 'Microsoft YaHei'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei'); styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor(0,0,0)

p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('WJ DSRC Protocol Test Desk 操作说明')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('适用版本 WJ_ComprehensiveTest_Desk_20260922').bold = True
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('文档范围：当前工程已实现的 DSRC 测试界面、台发设置、透传操作和测试结果展示')
doc.add_page_break()

doc.add_heading('1 使用范围与启动', level=1)
doc.add_paragraph('本说明面向使用 WJ_ComprehensiveTest_Desk_20260922.exe 进行 DSRC 协议测试的人员。程序当前提供三个一级产品 Tab，其中 DSRC 页面沿用现有测试界面和逻辑；OBU 串口和发卡页面为同级预留页面，当前仅显示后续添加提示。')
doc.add_paragraph('启动后先确认串口、台发设备和待测 OBU 已连接。未完成端口打开和初始化台发前，不应执行发送或测试按钮。')

doc.add_heading('2 一级产品 Tab', level=1)
doc.add_picture(str(ART / 'mainform_annotated.png'), width=Inches(6.8))
p = doc.add_paragraph('图 1  主窗体区域标注截图'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
table(doc, ['标号', '页面/区域', '当前状态与用途'], [
    ('1', 'WJ_DSRC_Protocol_test', '当前已实现的 DSRC 测试页面，包含台发通道、透传模式、标准测试和日常使用。'),
    ('2', '台发设置与状态', '选择串口、打开端口、配置台发参数、选择更新选项并执行初始化台发。'),
    ('3', '帧解析与日志区', '显示解析提示、测试过程信息和当前操作结果。'),
    ('4', '实时上下行帧区', '显示打开或保存的实时帧内容，并支持清空。'),
    ('5', '帧操作与测试区', '维护待发帧列表，选择交易帧/IC 卡/ESAM 操作并执行测试。'),
])
doc.add_heading('2.1 WJ_DSRC_Protocol_test', level=2)
doc.add_paragraph('选中该 Tab 后进入现有 DSRC 测试主页面。一级 Tab 仅用于页面切换，不会自动打开串口，也不会自动启动协议测试。')
doc.add_heading('2.2 WJ_OBUSerial_Desk', level=2)
doc.add_paragraph('当前为预留页面，页面显示“测试功能待后续添加”。当前版本不要依据该页面执行 OBU 串口测试。')
doc.add_heading('2.3 WJ_CreatCard_Desk', level=2)
doc.add_paragraph('当前为预留页面，页面显示“测试功能待后续添加”。当前版本不要依据该页面执行发卡测试。')

doc.add_heading('3 台发设置与通信端口', level=1)
doc.add_picture(str(ART / 'mainform_regions.png'), width=Inches(6.8))
p = doc.add_paragraph('图 2  台发设置和帧操作区域标注截图'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
table(doc, ['控件', '日常操作', '效果与注意事项'], [
    ('通信端口下拉框', '选择已识别的 COM 端口。', '决定台发通信使用的串口；端口未打开时不能进行设备交互。'),
    ('打开端口', '选择端口后点击。', '打开或关闭当前串口，并更新底部通信状态；失败时应先检查端口占用。'),
    ('BST 间隔', '输入毫秒数。', '控制 BST 发送间隔，必须使用合法数值。'),
    ('非 BST 帧间隔', '输入毫秒数。', '控制非 BST 帧之间的等待时间。'),
    ('帧发送次数', '输入重发次数。', '影响台发重发行为；过大可能延长测试时间。'),
    ('交易超时', '输入毫秒数。', '控制交易响应等待上限。'),
    ('功率等级/物理信道', '输入设备支持的参数。', '参与台发初始化，实际范围以设备和项目配置为准。'),
    ('BID/UnixTime/LLC/MAC 更新', '按测试要求勾选。', '影响初始化时相应字段是否动态更新；不要在未确认用例要求时随意改变。'),
    ('启用日志', '需要保存过程时勾选。', '启用测试过程日志记录，便于结果追溯。'),
    ('初始化台发', '参数确认后点击。', '将界面参数下发到台发设备；成功后方可进入测试。'),
])

doc.add_heading('4 二级页面与透传模式', level=1)
doc.add_heading('4.1 透传模式', level=2)
doc.add_paragraph('透传模式用于查看和维护 DSRC 帧。操作顺序通常为：打开串口 → 初始化台发 → 打开实时帧或加载帧内容 → 在帧操作区维护待发列表 → 点击测试。透传操作本身不等同于某个标准测试用例，是否通过以当前测试说明和设备返回为准。')
table(doc, ['区域/控件', '使用方式', '操作效果'], [
    ('帧解析', '点击上方实时上下行帧内容进行解析。', '在左侧区域显示帧解析结果和提示。'),
    ('实时上下行帧', '点击打开加载内容，点击保存保存当前内容，点击清空移除显示。', '仅维护显示内容和文件记录，不自动执行测试。'),
    ('帧操作列表', '选中条目后使用上移、下移、删除或清空。', '调整待发帧顺序或移除条目。'),
    ('交易帧/IC 卡/ESAM 选择器', '选择要加入待发列表的协议操作。', '为测试序列添加对应操作项。'),
    ('添加', '在对应选择器选中项目后点击。', '将操作加入帧操作列表。'),
    ('测试', '确认端口已打开、台发已初始化且列表顺序正确后点击。', '按当前列表向设备执行测试，并在日志/状态区域显示过程和结果。'),
    ('BUF 修改/清空 BUF', '需要维护缓冲区时使用。', '修改或清空 BUF 内容，影响后续相关帧操作。'),
])
doc.add_heading('4.2 日常使用', level=2)
doc.add_paragraph('日常使用页面用于项目已有的辅助操作。进入该页面后先阅读页面提示，再按照控件旁的说明执行。若操作需要设备响应，必须保留完整的发送帧、回复帧和状态信息。')
doc.add_heading('4.3 标准测试页面', level=2)
doc.add_paragraph('标准测试页面包含测试树、操作区和测试信息区。父级节点用于批量选择，叶子节点用于单项执行。双击用例用于查看流程说明；开始测试后应等待统一结果，不要重复点击测试入口。')

doc.add_heading('5 测试用例执行', level=1)
doc.add_paragraph('执行前确认：串口已打开、台发已初始化、待测设备处于可测试状态、测试树选择有效，并已确认当前算法、通道和日志选项。执行过程中观察测试信息区的步骤、帧摘要、失败原因和最终结论。')
table(doc, ['层级', '使用方式', '判定关注点'], [
    ('测试分组', '展开或勾选父级。', '父子勾选应保持一致，批量执行按显示顺序进行。'),
    ('子测试项', '展开后选择具体测试类型。', '确认进入的是目标协议流程，而不是同组的其他子项。'),
    ('叶子用例', '单独勾选并开始测试。', '按用例说明检查帧顺序、响应状态、超时或无响应等专用判定。'),
    ('停止', '测试进行中使用停止入口。', '请求取消后台流程，结束后恢复界面状态；设备已发送帧仍应保留在结果记录中。'),
])

doc.add_heading('6 日志与结果保存', level=1)
doc.add_paragraph('启用日志后，测试过程信息用于定位通信和判定问题。标准测试页面和北京地标测试页面按当前界面提供的结果保存选项保存测试文本和截图；未勾选时仅保留界面显示，不应将其误认为已生成结果文件。结果内容应包括测试名称、时间、请求/响应帧、DLL 返回码、判定和失败原因。')

doc.add_heading('7 异常处理与当前限制', level=1)
table(doc, ['现象', '处理建议'], [
    ('串口无法打开', '检查 COM 端口是否被其他程序占用，重新选择识别到的端口后再打开。'),
    ('初始化台发失败', '检查串口、设备连接和参数范围；初始化成功前不要开始测试。'),
    ('设备无响应或超时', '保存当前日志和完整帧，检查台发参数、物理信道、设备状态及用例的无响应判定。'),
    ('测试失败', '根据测试信息区确认失败步骤、期望状态、实际状态和完整回复帧，不要只依据弹窗文字判断。'),
    ('OBU 串口/发卡页面无功能', '当前版本属于预留页面，相关测试功能尚未添加。'),
])
doc.add_paragraph('当前版本未连接真实 RSU、OBU 或读卡器进行本说明生成过程中的实机回归；涉及空口响应、卡片数据、语音效果或永久写卡的行为，仍需在受控测试设备上确认。')

doc.save(OUT)
print(OUT)
